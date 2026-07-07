#!/usr/bin/env python3
"""Генератор presentation/Offline_iOS_Lecture.pptx.

Источник правды по СОСТАВУ и ПОРЯДКУ слайдов — метки `[СЛАЙД: …]` в
LECTURE_SPEECH.md (1 метка = 1 слайд). Содержимое слайдов задаётся здесь,
в структуре SLIDES (тезисы, таблицы, код — из LECTURE_PLAN.md и сути [РЕЧЬ]).

При рассинхроне заголовков/порядка/коммитов с речью скрипт падает с
понятной ошибкой — чтобы речь и колода не разъезжались.

Запуск:  .venv-pptx/bin/python scripts/generate_presentation.py
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SPEECH = ROOT / "LECTURE_SPEECH.md"
OUTPUT = ROOT / "presentation" / "Offline_iOS_Lecture.pptx"

# ---------------------------------------------------------------- оформление

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.55)
CONTENT_W = SLIDE_W - MARGIN_L - Inches(0.55)
TITLE_TOP = Inches(0.32)
BODY_TOP = Inches(1.42)
BODY_BOTTOM = Inches(7.08)
GAP = Inches(0.14)

C_BG = RGBColor(0x1A, 0x1A, 0x2E)        # фон
C_TITLE = RGBColor(0xFF, 0xFF, 0xFF)
C_BODY = RGBColor(0xE8, 0xE8, 0xF0)
C_ACCENT = RGBColor(0x4F, 0xC3, 0xF7)    # акцент
C_MUTED = RGBColor(0x9E, 0x9E, 0xB0)
C_CODE = RGBColor(0xA8, 0xE6, 0xCF)
C_PANEL = RGBColor(0x16, 0x21, 0x3E)     # плашки кода/таблиц
C_ROW = RGBColor(0x20, 0x20, 0x40)

FONT = "Helvetica Neue"
FONT_MONO = "Menlo"

MAX_CODE_LINES = 15
MAX_CODE_CHARS = 92

# ------------------------------------------------------------------- контент
#
# kind: "title" — титульный; "big" — крупная центральная типографика;
#       "content" — заголовок + блоки.
# blocks: ("bullets", [строки]) — префиксы: "! " акцент, "~ " приглушённый;
#         ("bullets18", [...]) — то же, крупнее (18pt);
#         ("key", "фраза") — крупный акцентный тезис;
#         ("table", {"headers": [...], "rows": [[...]], "widths": [...]});
#         ("code", "swift-код") / ("flow", "моноширинная схема");
#         ("demo", "описание live demo").

SLIDES: list[dict] = [
    {
        "title": "title — «О себе и том, что у меня болит»",
        "kind": "title",
        "block": "Блок 0 · Введение и проблема",
    },
    {
        "title": "Стартовая точка разработки",
        "kind": "content",
        "commit": "1",
        "block": "Блок 0 · Введение и проблема",
        "blocks": [
            ("demo", "Функционал приложения: лента, билет с QR, баланс → выключаем сеть → "
                     "белый экран / бесконечный спиннер / «Нет соединения, попробуйте позже»"),
            ("bullets", [
                "! «Попробуйте позже» — вежливый отказ: без сети приложение не умеет ничего показать",
                "Хотя данные уже могли лежать на телефоне: лента, билет, баланс",
                "Демо-бэкенд: JSONPlaceholder — баланс и билет ходят в один эндпоинт GET /users/1; "
                "у ответов честные Cache-Control и ETag",
                "Lorem Picsum — картинки к постам (seed по id поста); отдаёт 302-редирект на CDN — "
                "вернёмся к этому в кэше картинок",
                "Билет, баланс и QR собираются на клиенте из профиля: payload — в Ticket.qrPayload, "
                "QR рисуется на экране",
            ]),
        ],
    },
    {
        "title": "«Сеть — это ненадёжный внешний ресурс, а не данность»",
        "kind": "big",
        "block": "Блок 0 · Введение и проблема",
        "big": "Сеть — это ненадёжный внешний ресурс,\nа не данность",
        "sub": [
            "Метро · самолёт · лифт · роуминг · переполненный стадион · плохой Wi-Fi в кафе",
            "Сеть отваливается постоянно — это нормальное состояние, а не исключение",
            "Проектируем от «сеть обычно есть» → получаем приложение, которое «обычно работает»",
        ],
    },
    {
        "title": "Роудмап развития приложения",
        "kind": "content",
        "block": "Блок 0 · Введение и проблема",
        "blocks": [
            ("flow", "online-only   ──▶   кэш для чтения   ──▶   offline-first"),
            ("bullets", [
                "Отказоустойчивость — не галочка, а ступени: насколько приложение переживает пропадание сети",
                "online-only — без интернета белый экран или ошибка",
                "кэш для чтения — показываем уже загруженное, но не полагаемся как на данные пользователя",
                "offline-first — локальные данные = источник правды, сеть только синхронизирует; "
                "читаем и пишем офлайн",
                "~ План: блок 1 — кэш и инвалидация · блок 2 — offline-first, билет, outbox · "
                "блок 3 — безопасность, Keychain, FaceID · блок 4 — итоги, BLE, чек-лист",
            ]),
        ],
    },
    # -------------------------------------------------- Блок 1. Кэширование
    {
        "title": "«URLCache — почти бесплатный HTTP-кэш»",
        "kind": "content",
        "commit": "2",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "Прослойка между кодом и сетью: тот же session.data(for:), система сама решает — "
                "сеть или сохранённый ответ",
                "Memory — быстрый кэш в RAM, живёт пока запущено приложение; "
                "Disk — переживает перезапуск, лежит в Library/Caches",
                "! Library/Caches ОС вправе вычистить под нехватку места — вернёмся к этому ещё трижды",
                "Старые записи вытесняются автоматически при превышении ёмкости",
                "! Кэшем управляет не клиент — сервер, через HTTP-заголовки",
            ]),
            ("code",
             'let cache = URLCache(\n'
             '    memoryCapacity: 16 * 1024 * 1024,   // RAM: живёт пока запущено приложение\n'
             '    diskCapacity: 128 * 1024 * 1024,    // диск: Library/Caches\n'
             '    diskPath: "api-cache"\n'
             ')\n'
             'let config = URLSessionConfiguration.default\n'
             'config.urlCache = cache\n'
             'config.requestCachePolicy = .useProtocolCachePolicy'),
        ],
    },
    {
        "title": "«Два вопроса кэша: свежесть и валидация» + заголовки",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "1. «Ответ ещё свежий?» — да → из кэша мгновенно: ноль трафика, ноль задержки",
                "2. «Несвежий (stale) — можно дёшево перепроверить?» — да → условный запрос: "
                "«не менялось» или новая версия",
            ]),
            ("table", {
                "headers": ["Механизм", "Заголовки", "Смысл"],
                "widths": [2.2, 4.6, 5.4],
                "rows": [
                    ["Свежесть", "Cache-Control: max-age=N", "ответ свежий N секунд — в сеть не ходим"],
                    ["Свежесть", "no-store / no-cache", "не сохранять вообще / отдавать только после перепроверки"],
                    ["Валидация", "ETag ↔ If-None-Match", "«отпечаток» версии ресурса"],
                    ["Валидация", "Last-Modified ↔ If-Modified-Since", "то же по дате — проще, но грубее (1 с)"],
                    ["Ответ сервера", "304 Not Modified / 200 + тело", "«не менялось, бери из кэша» / новая версия"],
                ],
            }),
            ("bullets", [
                "~ Условные заголовки (If-None-Match и др.) URLSession подставляет сам — вручную писать не нужно",
            ]),
        ],
    },
    {
        "title": "«Стандарт HTTP (RFC 9110/9111)» + дерево решений URLCache",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "Это не фишка URLSession, а протокол HTTP: так же работают браузер, curl, OkHttp",
                "RFC 9110 «HTTP Semantics» + RFC 9111 «HTTP Caching» (2022)",
            ]),
            ("flow",
             "Есть ли в кэше ответ на этот запрос?\n"
             "├─ Нет → идём в сеть, сохраняем (если заголовки разрешают)\n"
             "└─ Да → он ещё свежий (max-age / Expires)?\n"
             "        ├─ Да → из кэша, в сеть НЕ идём\n"
             "        └─ Нет (stale) → есть ETag / Last-Modified?\n"
             "                ├─ Да → условный запрос → 304: тело из кэша · 200: обновляем\n"
             "                └─ Нет → полный запрос в сеть"),
            ("bullets", [
                "Пример: GET /posts → max-age=43200 + ETag. Через 5 минут — из кэша; "
                "через 13 часов — If-None-Match → 304 → тело из кэша",
                "! Сети нет совсем → URLSession бросает ошибку (stale, перепроверить не у кого) — "
                "нужен ручной офлайн-фолбэк",
            ]),
        ],
    },
    {
        "title": "requestCachePolicy — 4 политики",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("table", {
                "headers": ["Политика", "Поведение"],
                "widths": [5.0, 7.2],
                "rows": [
                    [".useProtocolCachePolicy", "по HTTP-заголовкам сервера — наш дефолт, пока есть сеть"],
                    [".reloadIgnoringLocalCacheData", "всегда сеть, кэш игнорируется"],
                    [".returnCacheDataElseLoad", "сначала кэш, сеть — если пусто"],
                    [".returnCacheDataDontLoad", "только кэш, в сеть не ходить"],
                ],
            }),
            ("code",
             'do {\n'
             '    let (data, response) = try await session.data(for: request)  // сеть + автокэш\n'
             '    return Fetched(value: try decode(data), dataSource: .network)\n'
             '} catch {\n'
             '    // офлайн: достаём последний ответ, даже если по HTTP он протух\n'
             '    if let cached = cache.cachedResponse(for: request),\n'
             '       let value = try? decode(cached.data) as T {\n'
             '        return Fetched(value: value, dataSource: .offlineCache)\n'
             '    }\n'
             '    throw APIError.transport(error)\n'
             '}'),
            ("bullets", [
                "«Данные от вчера» лучше белого экрана; источник помечаем через dataSource — "
                "к честной индикации в UI вернёмся в инвалидации",
            ]),
        ],
    },
    {
        "title": "«Картинки — отдельная история»",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets18", [
                "Картинок много и они тяжёлые",
                "Аватары, обложки, QR/штрихкоды билетов должны показываться мгновенно и офлайн",
                "! Базовый инструмент для памяти — NSCache: автоматически чистится под давлением памяти",
                "Это не Dictionary: система сама выкинет картинки, если приложению не хватает RAM — "
                "руками чистить не надо",
            ]),
        ],
    },
    {
        "title": "«Подводный камень с редиректами» + «URLCache: сохранили под B, ищем по A»",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "Многие CDN (и picsum.photos в демо) отдают картинку через 302 на финальный URL",
                "URLSession следует редиректу прозрачно: в коде — исходный адрес, внутри — два запроса",
                "! URLCache сохраняет ответ под финальным URL, а ищем мы по исходному → cache miss",
            ]),
            ("flow",
             "GET picsum.photos/seed/42/200/120 → 302 → GET fastly.../237.jpg → 200 + JPEG\n"
             "\n"
             "Сохранили под:  fastly.picsum.photos/.../237.jpg     JPEG на диске ✓\n"
             "Ищем по:        picsum.photos/seed/42/...            промах ✗"),
            ("bullets", [
                "JSON из GET /posts не страдает: 200 сразу, без редиректа",
                "NSCache в памяти работает — кладём сами под исходный URL; проблема после перезапуска "
                "или в офлайне, когда RAM пуста",
                "~ Аналогия: заказали на «склад A», курьер перенаправил на «склад B» — приходите на A: «у нас нет»",
            ]),
        ],
    },
    {
        "title": "решение — файловый кэш картинок",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "Для картинок отключаем URLCache и делаем свой файловый кэш: ключ = SHA256(исходный URL)",
                "Индексируем по адресу, который видит приложение, а не по тому, куда увёл редирект",
                "Три уровня: NSCache (RAM) → диск (Library/Caches) → сеть, после — кладём в оба",
                "! clearCache() чистит оба уровня разом — понадобится при логауте",
            ]),
            ("code",
             'final class ImageLoader {\n'
             '    private let memory = NSCache<NSURL, UIImage>()   // L1: RAM, авто-эвикция\n'
             '    let diskCacheURL: URL                            // L2: Library/Caches\n'
             '    func image(for url: URL) async throws -> UIImage {\n'
             '        if let hit = memory.object(forKey: url as NSURL) { return hit }     // L1\n'
             '        if let image = loadFromDisk(for: url) {                             // L2\n'
             '            memory.setObject(image, forKey: url as NSURL); return image\n'
             '        }\n'
             '        let (data, _) = try await session.data(from: url)                   // L3\n'
             '        let image = UIImage(data: data)!\n'
             '        saveToDisk(data, for: url)   // ключ файла = SHA256(исходного URL)\n'
             '        memory.setObject(image, forKey: url as NSURL)\n'
             '        return image\n'
             '    }\n'
             '}'),
        ],
    },
    {
        "title": "CacheManager — единая точка, три механизма",
        "kind": "content",
        "commit": "3",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("key", "Коммит 2 научил показывать офлайн. Коммит 3 — не показать протухшее "
                    "и не утечь чужие данные."),
            ("bullets", [
                "CacheManager — единая точка управления жизненным циклом кэша",
                "1. TTL — срок жизни записи задаёт приложение (в URLCache — сервер): "
                "второй, наш уровень контроля актуальности поверх HTTP",
                "2. Версия схемы — константа cacheSchemaVersion: сменили формат модели → "
                "весь старый кэш сброшен на старте. Одна строчка против класса багов после обновления",
                "3. Очистка по событию — логаут / смена пользователя (следующий слайд)",
                "! Паттерн cache-then-network: кэш мгновенно, свежее — фоном "
                "(по духу stale-while-revalidate)",
                "Честно сообщаем актуальность: DataSource → бейджи «данные могли устареть, обновляем», "
                "«оффлайн — сохранённые данные»",
            ]),
        ],
    },
    {
        "title": "🔑 «Очистка при логауте — это безопасность»",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("bullets", [
                "Классика: пользователь A залогинился → лента закэшировалась → A вышел → "
                "зашёл B → видит данные A",
                "Семейный телефон, рабочий планшет, демо-устройство — в проде находят постоянно",
                "! Единая точка: один clearAll() чистит все слои — HTTP-ответы, TTL-метаданные, "
                "картинки в памяти и на диске",
                "Самая частая ошибка: почистили URLCache, а картинки бывшего пользователя остались",
                "~ Позже в logout() добавятся Keychain и локальное хранилище — это архитектурная точка, "
                "а не разрозненные вызовы",
            ]),
            ("code",
             '// CacheManager — единая точка очистки всех слоёв кэша\n'
             'func clearAll() {\n'
             '    urlCache.removeAllCachedResponses()                        // HTTP-ответы\n'
             '    UserDefaults.standard.removeObject(forKey: timestampKey)   // TTL-метаданные\n'
             '    ImageLoader.shared.clearCache()                            // картинки: RAM + диск\n'
             '}\n'
             '\n'
             '// SessionStore.logout() — дёргает единую точку\n'
             'func logout() {\n'
             '    currentUser = nil\n'
             '    CacheManager.shared.clearAll()        // ← чужие данные не утекут\n'
             '}'),
        ],
    },
    {
        "title": "Итог блока 1 + мостик к блоку 2",
        "kind": "content",
        "block": "Блок 1 · Кэширование и инвалидация",
        "blocks": [
            ("key", "Кэш без стратегии инвалидации — это не оптимизация,\nа отложенный баг или утечка"),
            ("table", {
                "headers": ["Механизм", "Закрывает риск"],
                "widths": [4.2, 8.0],
                "rows": [
                    ["TTL", "протухшие данные"],
                    ["Версия схемы", "несовместимый формат после обновления из App Store"],
                    ["Очистка по событию (логаут)", "чужие данные у пользователя B"],
                ],
            }),
            ("bullets", [
                "Уберите любой из трёх — получите конкретный класс багов. Этот подраздел не режем никогда",
                "! Мостик: у некоторых данных свой срок годности, не совпадающий со сроком жизни кэша — "
                "платёжный токен, билет, баланс",
            ]),
            ("flow", "кэш (производное, не жалко потерять)   ──▶   источник правды (терять нельзя) — блок 2"),
        ],
    },
    # ---------------------------------------------------- Блок 2. Offline-first
    {
        "title": "«Кэш vs источник правды»",
        "kind": "content",
        "commit": "4",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("bullets", [
                "! Переворачиваем парадигму: локальное хранилище — источник правды, "
                "сеть лишь синхронизирует его",
            ]),
            ("table", {
                "headers": ["", "Кэш (коммиты 2–3)", "Источник правды (коммит 4)"],
                "widths": [2.0, 5.0, 5.2],
                "rows": [
                    ["Роль", "фолбэк «если сеть упала»", "первоисточник, читается всегда"],
                    ["Природа данных", "производные — можно потерять", "данные пользователя — терять нельзя"],
                    ["Папка", "Library/Caches — ОС может удалить", "Application Support — ОС не трогает, бэкапится"],
                    ["Поток", "сеть → показать; кэш при ошибке", "диск → показать сразу; сеть обновляет фоном"],
                ],
            }),
            ("bullets", [
                "! Для билета фолбэк-кэш недопустим: ОС вычистит Caches под нехватку места — "
                "пассажир останется без QR на стойке",
            ]),
        ],
    },
    {
        "title": "папки песочницы",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("table", {
                "headers": ["Папка", "Назначение", "ОС удаляет?", "В бэкапе?"],
                "widths": [3.4, 5.2, 2.0, 1.6],
                "rows": [
                    ["Documents", "пользовательские файлы, видны в Finder", "нет", "да"],
                    ["Application Support", "данные приложения — билет здесь", "нет", "да"],
                    ["Library/Caches", "URLCache, картинки — можно перекачать", "да", "нет"],
                    ["tmp", "временные файлы", "да", "нет"],
                ],
            }),
            ("bullets", [
                "! Правило: можно перекачать из сети → Caches; нельзя потерять → Application Support / Documents",
                "Один билет → простой generic-стор: Codable + JSON-файл; много данных → "
                "SQLite / Core Data / SwiftData",
            ]),
            ("code",
             'final class LocalStore {   // Codable + файлы в Application Support\n'
             '    func save<T: Encodable>(_ value: T, forKey key: String) {\n'
             '        guard let data = try? JSONEncoder().encode(value) else { return }\n'
             '        try? data.write(to: fileURL(key), options: .atomic)\n'
             '    }\n'
             '    func load<T: Decodable>(forKey key: String) -> T? {\n'
             '        guard let data = try? Data(contentsOf: fileURL(key)) else { return nil }\n'
             '        return try? JSONDecoder().decode(T.self, from: data)\n'
             '    }\n'
             '}'),
        ],
    },
    {
        "title": "local-first поток билета",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("bullets", [
                "loadLocal — билет с диска мгновенно: 0 мс, работает в авиарежиме",
                "refreshFromNetwork — фоном тихо обновляем источник правды, экран не блокируем",
                "Статус синхронизации: «синхронизировано» / «обновляем» / «офлайн» — "
                "честность из блока 1 продолжается",
                "! Экран ошибки — только при самом первом запуске без сети; "
                "во всех остальных сценариях билет показывается всегда",
            ]),
            ("code",
             'final class TicketRepository {\n'
             '    // ① мгновенно с диска — работает в авиарежиме, 0 мс\n'
             '    func loadLocal() -> Ticket? { store.load(forKey: ticketKey) }\n'
             '\n'
             '    // ② в фоне обновляем ИСТОЧНИК ПРАВДЫ, не блокируя UI\n'
             '    func refreshFromNetwork() async throws -> Ticket {\n'
             '        let result: Fetched<RemoteUser> = try await client.get(.user(id: 1))\n'
             '        let ticket = Ticket(user: result.value)\n'
             '        store.save(ticket, forKey: ticketKey)\n'
             '        return ticket\n'
             '    }\n'
             '\n'
             '    // ③ при логауте персональные данные не остаются на устройстве\n'
             '    func clearLocal() { store.remove(forKey: ticketKey) }\n'
             '}'),
        ],
    },
    {
        "title": "«QR строится из локального payload»",
        "kind": "big",
        "block": "Блок 2 · Offline-first",
        "big": "QR-код строится из локально сохранённого payload,\nа не из живого ответа сервера",
        "sub": [
            "Посадку можно пройти полностью офлайн",
            "Сеть нужна лишь для того, чтобы держать локальную копию свежей",
            "До этого билет жил в памяти ViewModel и «случайно» подхватывался из URLCache — "
            "на кэш полагаться нельзя, его вычистят",
        ],
    },
    {
        "title": "Итог: кэш vs источник правды",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("key", "Кэш: «что показать, если сеть упала»"),
            ("key", "Источник правды: «данные всегда здесь, локально;\nсеть — лишь механизм обновления»"),
            ("bullets18", [
                "Для билета, баланса, черновиков — только второй подход",
                "И лежать они должны в Application Support, а не в Caches",
                "~ Промежуточная сверка: дальше — offline не только для чтения, но и для записи",
            ]),
        ],
    },
    {
        "title": "Outbox — идея и компоненты",
        "kind": "content",
        "commit": "5",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("bullets", [
                "Пользователь не только читает — он действует: лайк, отметка, форма. "
                "В наивном приложении действие = немедленный запрос, без сети оно теряется",
                "! Outbox: действие сначала пишется в локальную очередь, UI реагирует оптимистично, "
                "доставкой занимается фоновый воркер",
            ]),
            ("table", {
                "headers": ["Компонент", "Роль"],
                "widths": [3.4, 8.8],
                "rows": [
                    ["NetworkMonitor", "обёртка над NWPathMonitor — знает о появлении сети заранее, до попытки запроса"],
                    ["OutboxItem", "clientID (UUID) + действие + статус + счётчик ретраев"],
                    ["OutboxStore", "персистентная очередь (outbox.json) — переживает перезапуск"],
                    ["OutboxProcessor", "воркер: pending → inFlight → отправка, ретраи с backoff"],
                ],
            }),
            ("flow", "тап (офлайн) → enqueue → диск (pending) → появилась сеть → processQueue → POST → отправлено"),
        ],
    },
    {
        "title": "🔑 идемпотентность",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("bullets", [
                "clientID — UUID: генерируется один раз и НЕ меняется между ретраями; "
                "летит в заголовке Idempotency-Key",
                "Сценарий: лайк отправлен, сервер применил, ответ потерялся → клиент думает «не дошло» и повторяет",
                "! Без ключа — два лайка или две оплаты. С ключом сервер видит «этот clientID уже обрабатывал» "
                "и не задваивает",
            ]),
            ("code",
             '// clientID генерируется ОДИН раз при постановке в очередь…\n'
             'let item = OutboxItem(clientID: UUID(), action: action, retryCount: 0)\n'
             '\n'
             '// …и не меняется между ретраями\n'
             'let _: PostResponse = try await api.post(\n'
             '    .posts, body: payload,\n'
             '    idempotencyKey: item.clientID.uuidString   // → заголовок Idempotency-Key\n'
             ')'),
            ("bullets", [
                "~ В демо JSONPlaceholder ключ не дедуплицирует — реализована клиентская половина контракта; "
                "серверную обеспечивает бэкенд",
            ]),
        ],
    },
    {
        "title": "сценарий демо outbox",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("demo", "Видео/live: outbox переживает офлайн и перезапуск приложения"),
            ("bullets18", [
                "1. Airplane Mode → тап «лайк» → в Настройках «в очереди: 1»",
                "2. Убить и перезапустить приложение (всё ещё офлайн) → очередь на месте — прочитана с диска",
                "3. Выключить Airplane Mode → NWPathMonitor ловит переход → воркер сам отправляет → "
                "очередь пустеет",
                "! Пользователь не сделал ничего особенного — всё произошло само",
            ]),
        ],
    },
    {
        "title": "доп. offline-кейсы — одним слайдом",
        "kind": "content",
        "block": "Блок 2 · Offline-first",
        "blocks": [
            ("bullets18", [
                "Карты и маршруты офлайн — заранее скачанный регион",
                "Штрихкод карты лояльности на кассе",
                "Контент «скачать для офлайна»: книги, подкасты, курсы — с синхронизацией прогресса",
                "Черновики, переживающие закрытие приложения",
                "! Везде одни и те же три идеи: локальный источник правды · "
                "честная индикация актуальности · очередь на запись",
            ]),
        ],
    },
    # ------------------------------------------------- Блок 3. Безопасность
    {
        "title": "«Без модели угроз советы = карго-культ»",
        "kind": "big",
        "block": "Блок 3 · Безопасность и хранение",
        "big": "Без модели угроз любые советы\nпо безопасности — карго-культ",
        "sub": [
            "«Шифруйте всё» звучит солидно, но не отвечает на вопрос «от кого»",
            "Нельзя «защититься вообще» — можно защититься от конкретного противника",
            "Весь блок построен от угроз, а не от списка API",
        ],
    },
    {
        "title": "от кого защищаемся",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("bullets18", [
                "Другое приложение — песочница iOS изолирует, но ошибки в хранении бьют по нам",
                "Вор телефона — физический доступ к устройству",
                "Бэкап в iCloud / iTunes — секреты могут уехать на другое устройство",
                "Jailbreak — ослабленная песочница",
                "Чужие глаза — коллега, ребёнок, камера над плечом",
                "Из блока 1 — компрометация канала и утечка данных между пользователями после логаута",
                "! Под каждый сценарий — свой инструмент, а не «шифруйте всё»",
            ]),
        ],
    },
    {
        "title": "карта хранилищ",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("table", {
                "headers": ["Хранилище", "Для чего", "Чего нельзя"],
                "widths": [3.6, 4.6, 4.0],
                "rows": [
                    ["Keychain", "токены, пароли, ключи шифрования", "большие объёмы данных"],
                    ["UserDefaults", "настройки, флаги", "секреты/токены — это plain-text plist"],
                    ["Documents / App Support", "пользовательские данные, бэкапятся", "временный кэш"],
                    ["Caches / tmp", "то, что можно перекачать", "то, что нельзя потерять"],
                    ["Core Data / SQLite", "офлайн-датасеты", "секреты без шифрования"],
                ],
            }),
            ("bullets", [
                "! UserDefaults — открытый текст: физически это .plist в Library/Preferences",
                "Его видно в бэкапе, его открывает текстовый редактор. Любой токен там — секрет, отданный наружу",
            ]),
        ],
    },
    {
        "title": "папки песочницы + бэкап",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("table", {
                "headers": ["Папка", "В бэкапе iCloud", "ОС может удалить"],
                "widths": [4.2, 4.0, 4.0],
                "rows": [
                    ["Documents", "да", "нет"],
                    ["Application Support", "да", "нет"],
                    ["Library/Caches", "нет", "да — под нехватку места"],
                    ["tmp", "нет", "да — в любой момент"],
                ],
            }),
            ("bullets", [
                "isExcludedFromBackup — исключить конкретный файл из бэкапа "
                "(например, большой скачанный датасет)",
                "! Выбор папки — это и надёжность, и безопасность: билеты и черновики ≠ кэш, "
                "они должны лежать в правильном месте",
            ]),
        ],
    },
    {
        "title": "Keychain — зачем и как",
        "kind": "content",
        "commit": "6",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("bullets", [
                "Первый настоящий секрет — access token после логина. По карте хранилищ место одно — Keychain",
                "Зашифрованное хранилище ОС: AES-256, ключи живут в Secure Enclave и не покидают его",
                "API старый и сишный (SecItemAdd, SecItemCopyMatching) → обёртка KeychainStore: "
                "save / read / delete",
                "Логин → save; старт приложения → read (сессия переживает перезапуск); логаут → delete",
                "~ Для маленьких секретов — токенов и ключей, не для больших объёмов данных",
            ]),
            ("code",
             'final class KeychainStore {\n'
             '    func save(_ data: Data, account: String) throws {\n'
             '        var query: [CFString: Any] = [\n'
             '            kSecClass: kSecClassGenericPassword,   // класс для app-токенов\n'
             '            kSecAttrService: service, kSecAttrAccount: account,\n'
             '            kSecValueData: data,\n'
             '            kSecAttrAccessible: kSecAttrAccessibleAfterFirstUnlockThisDeviceOnly]\n'
             '        let status = SecItemAdd(query as CFDictionary, nil)\n'
             '        guard status == errSecSuccess else { throw KeychainError.unexpectedStatus(status) }\n'
             '    }\n'
             '    // + read → SecItemCopyMatching, delete → SecItemDelete\n'
             '}'),
        ],
    },
    {
        "title": "классы доступности",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("table", {
                "headers": ["Класс доступности", "Когда ОС отдаёт секрет"],
                "widths": [4.6, 7.6],
                "rows": [
                    ["WhenUnlocked", "только при разблокированном экране"],
                    ["AfterFirstUnlock", "после первой разблокировки после включения — доступен и в фоне"],
                    ["+ суффикс ThisDeviceOnly", "секрет не попадает в бэкап и не переезжает на другое устройство"],
                ],
            }),
            ("bullets", [
                "! Наш выбор — AfterFirstUnlockThisDeviceOnly, и оба слова осознанные:",
                "AfterFirstUnlock — outbox-воркеру из блока 2 нужно слать запросы с токеном в фоне, "
                "без участия пользователя",
                "ThisDeviceOnly — в модели угроз есть «бэкап в iCloud»: "
                "компрометация бэкапа не должна означать компрометацию токена",
            ]),
        ],
    },
    {
        "title": "контраст UserDefaults vs Keychain",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("demo", "Кнопка пишет токен в UserDefaults → открываем Library/Preferences/<bundle>.plist — "
                     "токен виден глазами. Рядом токен в Keychain — так просто не достать"),
            ("table", {
                "headers": ["", "UserDefaults", "Keychain"],
                "widths": [2.2, 5.0, 5.0],
                "rows": [
                    ["Физически", "plain-text .plist в Library/Preferences", "зашифрованная БД ОС"],
                    ["Шифрование", "нет", "AES-256, ключи в Secure Enclave"],
                    ["Бэкап", "уезжает в iCloud как есть", "с ThisDeviceOnly — не уезжает"],
                    ["Что класть", "настройки, флаги, имя пользователя", "секреты: токены, пароли, ключи"],
                ],
            }),
            ("bullets", [
                "Имя пользователя намеренно в UserDefaults — это не секрет. Токен — только в Keychain",
                "! Безопасность — не «всё шифровать», а «класть каждый тип данных в правильное место»",
                "~ Файлы шифрует Data Protection при блокировке, ключи — в Secure Enclave; "
                "Keychain использует эту машинерию за вас",
            ]),
        ],
    },
    {
        "title": "FaceID — гейт, а не шифрование",
        "kind": "content",
        "commit": "7",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("key", "FaceID — это гейт доступа, а не шифрование"),
            ("bullets", [
                "Данные билета уже на диске и зашифрованы Data Protection — FaceID их не расшифровывает",
                "Он решает один вопрос: показать пользователю или нет",
                "! Замок на двери, не сейф. Сейф — Secure Enclave и Keychain",
            ]),
            ("table", {
                "headers": ["Политика LAPolicy", "Что разрешает", "Где используем"],
                "widths": [5.2, 3.8, 3.2],
                "rows": [
                    [".deviceOwnerAuthenticationWithBiometrics", "только FaceID / TouchID", "гейт на билет"],
                    [".deviceOwnerAuthentication", "биометрия или пароль устройства", "фолбэк (симулятор без биометрии)"],
                ],
            }),
        ],
    },
    {
        "title": "гейт на TicketView — реблокировка при фоне + SecAccessControl",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("bullets", [
                "Приложение ушло в фон → экран билета снова блокируется: гейт, "
                "а не «разблокировал один раз навсегда»",
                "! Платёжный токен привязан к биометрии на уровне Keychain: SecAccessControl(.biometryCurrentSet)",
                "Приложение не вызывает проверку само — биометрию запрашивает Keychain через Secure Enclave; "
                "секрет физически не отдаётся без живого FaceID",
                ".biometryAny: добавили новое лицо — секрет всё ещё читается. "
                ".biometryCurrentSet: секрет недоступен — защита от подмены. Для платёжных данных — только CurrentSet",
            ]),
            ("code",
             'let access = SecAccessControlCreateWithFlags(\n'
             '    kCFAllocatorDefault,\n'
             '    kSecAttrAccessibleWhenUnlockedThisDeviceOnly,\n'
             '    .biometryCurrentSet,        // 🔑 не .biometryAny\n'
             '    &error\n'
             ')\n'
             'query[kSecAttrAccessControl] = access   // секрет привязан к текущей биометрии\n'
             '\n'
             '// чтение: FaceID-промпт показывает сам Keychain через Secure Enclave\n'
             'query[kSecUseAuthenticationContext] = LAContext()\n'
             'SecItemCopyMatching(query as CFDictionary, &result)'),
        ],
    },
    {
        "title": "два токена",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("table", {
                "headers": ["Токен", "Класс / ACL", "Чтение"],
                "widths": [3.2, 4.8, 4.2],
                "rows": [
                    ["accessToken (коммит 6)", "AfterFirstUnlockThisDeviceOnly", "без биометрии — нужен фоновому outbox"],
                    ["paymentToken (коммит 7)", "SecAccessControl(.biometryCurrentSet)", "только через FaceID на уровне ОС"],
                ],
            }),
            ("bullets", [
                "Два секрета — две политики, потому что требования разные: фоновые запросы vs платёж",
                "! logout удаляет оба — в той же единой точке очистки, что мы завели в блоке 1",
                "Теперь в ней пять слоёв: Keychain ×2 · кэш · локальное хранилище · outbox",
            ]),
        ],
    },
    {
        "title": "не забыть Info.plist",
        "kind": "content",
        "block": "Блок 3 · Безопасность и хранение",
        "blocks": [
            ("key", "NSFaceIDUsageDescription — обязательный ключ в Info.plist"),
            ("bullets18", [
                "! Без него обращение к FaceID не «вернёт ошибку», а крашит приложение",
                "Маленькое, но злое — проверьте в каждом проекте с биометрией",
            ]),
            ("code",
             '<key>NSFaceIDUsageDescription</key>\n'
             '<string>Подтверждаем личность перед показом билета и оплатой</string>'),
        ],
    },
    # ------------------------------------------------------- Блок 4. Итоги
    {
        "title": "«Оффлайн-оплата по BLE» — на словах",
        "kind": "content",
        "block": "Блок 4 · Итоги",
        "blocks": [
            ("bullets", [
                "Самый эффектный кейс — собирает всё вместе: оплата телефоном у терминала "
                "вообще без интернета (как транспортные карты)",
                "Пока есть сеть — приложение заранее загружает пачку подписанных сервером платёжных токенов",
                "У каждого токена: срок жизни, лимит по сумме и количеству",
                "Терминал офлайн проверяет подпись сервера — в сеть не выходит",
                "Токены лежат в Keychain под .biometryCurrentSet — оплата подтверждается FaceID",
                "Транзакции уходят через outbox и досинхронизируются, когда сеть появится",
                "! Зачем ограничения: украденный телефон ≠ бесконечная оффлайн-оплата — "
                "токенов мало, они протухнут и ограничены по сумме",
                "~ Мостик из блока 1 в полный рост: у платёжного токена свой срок годности — "
                "кэш не подходит, нужен источник правды с явной валидацией",
            ]),
        ],
    },
    {
        "title": "типичные ошибки",
        "kind": "content",
        "block": "Блок 4 · Итоги",
        "blocks": [
            ("bullets18", [
                "✗ Показывают протухший кэш без индикации",
                "✗ Кладут токены в UserDefaults",
                "✗ Хранят билеты в Library/Caches",
                "✗ FaceID только в UI, без привязки к Keychain",
                "✗ Забывают чистить всё при логауте",
                "✗ Считают «попробуйте позже» нормальным офлайном",
            ]),
        ],
    },
    {
        "title": "чек-лист отказоустойчивого приложения",
        "kind": "content",
        "block": "Блок 4 · Итоги",
        "blocks": [
            ("bullets18", [
                "✓ Локальный источник правды; сеть только синхронизирует",
                "✓ У кэша — TTL, версия схемы, очистка при логауте",
                "✓ Пользователю честно показываем актуальность данных",
                "✓ Запись без сети — в outbox с идемпотентностью",
                "✓ Секреты — только Keychain с правильным классом доступности",
                "✓ Файлы — в правильных папках; критичное исключено из ненужного бэкапа",
                "✓ Чувствительное — за FaceID; платёжные токены привязаны к биометрии",
                "✓ Офлайн-токены оплаты — срок, лимит, подпись",
                "~ Можно сфотографировать",
            ]),
        ],
    },
    {
        "title": "финальная мысль",
        "kind": "big",
        "block": "Блок 4 · Итоги",
        "big": "Сеть ненадёжна. Данные живут локально.\nПользователю всегда говорим правду\nо том, что он видит.",
        "sub": [
            "Отказоустойчивость — не фича, которую прикручивают в конце, а способ мышления",
            "Постройте приложение так — и авиарежим перестанет быть проблемой",
            "Спасибо! Вопросы",
        ],
    },
]

# --------------------------------------------------- сверка с LECTURE_SPEECH

SLIDE_RE = re.compile(r"`\[СЛАЙД:\s*(.+?)\]`")
COMMIT_RE = re.compile(r"\[КОММИТ\s*(\d+)\]")


def speech_slides(text: str) -> list[tuple[str, str | None]]:
    """[(заголовок, коммит|None)] — по меткам речи; легенда (строки '>') пропускается."""
    result = []
    for line in text.splitlines():
        if line.lstrip().startswith(">"):
            continue
        m = SLIDE_RE.search(line)
        if m:
            cm = COMMIT_RE.search(line)
            result.append((m.group(1).strip(), cm.group(1) if cm else None))
    return result


def verify_against_speech() -> None:
    labels = speech_slides(SPEECH.read_text(encoding="utf-8"))
    expected = [(s["title"], s.get("commit")) for s in SLIDES]
    if labels == expected:
        return
    msgs = [f"Рассинхрон речи и колоды: в речи {len(labels)} меток, в скрипте {len(expected)} слайдов."]
    for i in range(max(len(labels), len(expected))):
        a = labels[i] if i < len(labels) else ("<нет метки>", None)
        b = expected[i] if i < len(expected) else ("<нет слайда>", None)
        if a != b:
            msgs.append(f"  слайд {i + 1}: речь = {a!r}  /  скрипт = {b!r}")
    raise SystemExit("\n".join(msgs))


def check_code_limits() -> None:
    for s in SLIDES:
        for kind, payload in s.get("blocks", []):
            if kind != "code":
                continue
            lines = payload.splitlines()
            assert len(lines) <= MAX_CODE_LINES, f"{s['title']}: код длиннее {MAX_CODE_LINES} строк"
            for ln in lines:
                assert len(ln) <= MAX_CODE_CHARS, f"{s['title']}: строка кода длиннее {MAX_CODE_CHARS}: {ln!r}"


# ----------------------------------------------------------------- рендеринг


def display_title(title: str) -> str:
    """Заголовок для показа: как в речи, но с заглавной первой буквы."""
    return title[0].upper() + title[1:] if title else title


def est_lines(text: str, chars: int) -> int:
    return max(1, math.ceil(len(text) / chars))


def block_height(block) -> float:
    kind, payload = block
    if kind in ("bullets", "bullets18"):
        chars = 105 if kind == "bullets" else 95
        per = 0.30 if kind == "bullets" else 0.34
        return sum(est_lines(t.lstrip("!~ "), chars) * per + 0.07 for t in payload)
    if kind == "key":
        return est_lines(payload, 80) * 0.38 + 0.1 + 0.34 * payload.count("\n")
    if kind == "table":
        return 0.37 * (len(payload["rows"]) + 1) + 0.06
    if kind in ("code", "flow"):
        return 0.205 * len(payload.splitlines()) + 0.28
    if kind == "demo":
        return est_lines(payload, 90) * 0.28 + 0.52
    raise ValueError(kind)


def set_bg(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def para(tf, text, *, first=False, size=16, color=C_BODY, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, space_after=4):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    for run in p.runs or []:
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font
    return p


def add_box(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def render_bullets(slide, top, items, size) -> None:
    tf = add_box(slide, MARGIN_L, top, CONTENT_W, Inches(0.5))
    first = True
    for raw in items:
        if raw.startswith("! "):
            text, color, bold = "• " + raw[2:], C_ACCENT, True
        elif raw.startswith("~ "):
            text, color, bold = "• " + raw[2:], C_MUTED, False
        else:
            text, color, bold = "• " + raw, C_BODY, False
        para(tf, text, first=first, size=size, color=color, bold=bold, space_after=5)
        first = False


def render_key(slide, top, text) -> None:
    tf = add_box(slide, MARGIN_L, top, CONTENT_W, Inches(0.6))
    first = True
    for line in text.split("\n"):
        para(tf, line, first=first, size=21, color=C_ACCENT, bold=True, space_after=2)
        first = False


def render_table(slide, top, spec) -> None:
    headers, rows = spec["headers"], spec["rows"]
    n_rows, n_cols = len(rows) + 1, len(headers)
    height = Inches(0.37 * n_rows)
    frame = slide.shapes.add_table(n_rows, n_cols, MARGIN_L, top, CONTENT_W, height)
    table = frame.table
    table.first_row = False
    table.horz_banding = False

    weights = spec.get("widths") or [1] * n_cols
    total = sum(weights)
    for i, w in enumerate(weights):
        table.columns[i].width = Emu(int(CONTENT_W * w / total))
    for r in range(n_rows):
        table.rows[r].height = Inches(0.37)

    for c, text in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_PANEL
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.text_frame.word_wrap = True
        para(cell.text_frame, text, first=True, size=13, color=C_ACCENT, bold=True, space_after=0)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_ROW
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.text_frame.word_wrap = True
            mono = text.startswith(".") or text.startswith(("kSec", "Sec", "URL", "NS", "ETag", "Cache-", "Last-", "304", "Idempotency"))
            para(cell.text_frame, text, first=True, size=13,
                 color=C_BODY if c else C_TITLE, bold=(c == 0),
                 font=FONT_MONO if mono else FONT, space_after=0)


def render_mono_panel(slide, top, height, text, color) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN_L, top, CONTENT_W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_PANEL
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    first = True
    for line in text.splitlines():
        para(tf, line if line else " ", first=first, size=12, color=color, font=FONT_MONO, space_after=1)
        first = False


def render_demo(slide, top, height, text) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN_L, top, CONTENT_W, height)
    shape.fill.background()
    shape.line.color.rgb = C_ACCENT
    shape.line.width = Pt(1.5)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    para(tf, "LIVE DEMO", first=True, size=14, color=C_ACCENT, bold=True, space_after=2)
    para(tf, text, size=14, color=C_BODY, space_after=0)


def render_footer(slide, idx, total, block) -> None:
    tf = add_box(slide, MARGIN_L, Inches(7.12), Inches(7), Inches(0.3))
    para(tf, block, first=True, size=10, color=C_MUTED, space_after=0)
    tf = add_box(slide, Inches(12.0), Inches(7.12), Inches(0.8), Inches(0.3))
    para(tf, f"{idx} / {total}", first=True, size=10, color=C_MUTED,
         align=PP_ALIGN.RIGHT, space_after=0)


def render_title_header(slide, data) -> None:
    title = display_title(data["title"])
    size = 28 if len(title) < 52 else 24 if len(title) < 72 else 20
    tf = add_box(slide, MARGIN_L, TITLE_TOP, CONTENT_W - Inches(1.4), Inches(1.0))
    para(tf, title, first=True, size=size, color=C_TITLE, bold=True, space_after=0)
    if data.get("commit"):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.55), Inches(0.38), Inches(1.25), Inches(0.38))
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_PANEL
        shape.line.color.rgb = C_ACCENT
        shape.line.width = Pt(1)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, f"commit {data['commit']}", first=True, size=12, color=C_ACCENT,
             bold=True, font=FONT_MONO, align=PP_ALIGN.CENTER, space_after=0)


def render_title_slide(slide) -> None:
    tf = add_box(slide, MARGIN_L, Inches(1.7), CONTENT_W, Inches(1.8))
    para(tf, "Превращаем онлайн-приложение", first=True, size=38, color=C_TITLE, bold=True, space_after=0)
    para(tf, "в отказоустойчивое", size=38, color=C_TITLE, bold=True, space_after=10)
    para(tf, "Offline-режим и безопасность данных в iOS", size=22, color=C_ACCENT, space_after=0)
    tf = add_box(slide, MARGIN_L, Inches(4.35), CONTENT_W, Inches(1.9))
    para(tf, "Илья · мобильный разработчик в Финтехе — бесконтактные платежи и цифровой рубль",
         first=True, size=17, color=C_BODY, space_after=8)
    para(tf, "Что болит: очередь на концерте, сканер билетов — один человек в минуту, "
             "и даже пиво без сети не купить", size=15, color=C_MUTED, space_after=8)
    para(tf, "90 минут · 3 ступени отказоустойчивости · 7 коммитов демо-приложения",
         size=15, color=C_MUTED, space_after=0)


def render_big_slide(slide, data) -> None:
    lines = data["big"].split("\n")
    tf = add_box(slide, MARGIN_L, Inches(2.0), CONTENT_W, Inches(2.2))
    first = True
    for line in lines:
        para(tf, line, first=first, size=32, color=C_ACCENT, bold=True,
             align=PP_ALIGN.CENTER, space_after=4)
        first = False
    tf = add_box(slide, Inches(1.1), Inches(2.35 + 0.55 * len(lines)), SLIDE_W - Inches(2.2), Inches(2.4))
    first = True
    for line in data.get("sub", []):
        para(tf, line, first=first, size=18, color=C_BODY, align=PP_ALIGN.CENTER, space_after=10)
        first = False


def render_slide(prs, data, idx, total) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_BG)

    if data["kind"] == "title":
        render_title_slide(slide)
        render_footer(slide, idx, total, data["block"])
        return
    if data["kind"] == "big":
        render_big_slide(slide, data)
        render_footer(slide, idx, total, data["block"])
        return

    render_title_header(slide, data)

    blocks = data["blocks"]
    heights = [block_height(b) for b in blocks]
    avail = (BODY_BOTTOM - BODY_TOP) / Inches(1)
    need = sum(heights) + (len(blocks) - 1) * GAP / Inches(1)
    assert need <= avail + 0.01, f"{data['title']}: контент не помещается ({need:.2f} > {avail:.2f} in)"

    top = BODY_TOP
    for block, h in zip(blocks, heights):
        kind, payload = block
        hh = Inches(h)
        if kind == "bullets":
            render_bullets(slide, top, payload, size=16)
        elif kind == "bullets18":
            render_bullets(slide, top, payload, size=18)
        elif kind == "key":
            render_key(slide, top, payload)
        elif kind == "table":
            render_table(slide, top, payload)
        elif kind == "code":
            render_mono_panel(slide, top, hh, payload, C_CODE)
        elif kind == "flow":
            render_mono_panel(slide, top, hh, payload, C_ACCENT)
        elif kind == "demo":
            render_demo(slide, top, hh, payload)
        top = top + hh + GAP

    render_footer(slide, idx, total, data["block"])


def main() -> None:
    verify_against_speech()
    check_code_limits()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = len(SLIDES)
    for i, data in enumerate(SLIDES, start=1):
        render_slide(prs, data, i, total)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"OK: {total} слайдов → {OUTPUT.relative_to(ROOT)} (сверено с {SPEECH.name})")


if __name__ == "__main__":
    main()
